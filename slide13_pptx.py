"""
Slide 13 - Headline Result : native, EDITABLE PowerPoint (16:9).
Same dark/coral design and the SAME real report values as the PNG mockup.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

BG     = RGBColor.from_string("0E1117")
PANEL  = RGBColor.from_string("171C26")
CORAL  = RGBColor.from_string("FF6F61")
CORALD = RGBColor.from_string("E85A4F")
INK    = RGBColor.from_string("F2F4F8")
MUTE   = RGBColor.from_string("9AA4B2")
GRID   = RGBColor.from_string("2A323F")
HILITE = RGBColor.from_string("241A18")
FONT   = "Calibri"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank


def rect(x, y, w, h, fill, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def text(x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for (s, size, color, bold) in line:
            r = p.add_run(); r.text = s
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.name = FONT
    return tb

# background + top rule
rect(0, 0, 13.333, 7.5, BG)
rect(0, 0, 13.333, 0.10, CORAL)

# kicker + headline
text(0.6, 0.30, 6.5, 0.4, [[("HEADLINE RESULT", 16, CORAL, True)]])
text(0.6, 0.68, 6.6, 1.4,
     [[("10% of the labels.", 36, INK, True)],
      [("93% of the Dice.",  36, INK, True)]])

# giant 10x
text(0.3, 1.75, 6.2, 2.5, [[("10×", 200, CORAL, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(0.3, 4.25, 6.2, 0.8,
     [[("FEWER LABELS", 23, INK, True)],
      [("10%  vs  100%  annotation budget", 13, MUTE, False)]],
     align=PP_ALIGN.CENTER)

# chips
def chip(y, icon, title, big, small):
    x, w, h = 7.45, 5.3, 1.02
    rect(x, y, w, h, PANEL, line=GRID, lw=1.25,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(x+0.06, y+0.12, 0.10, h-0.24, CORAL)                      # accent bar
    rect(x+0.30, y+h/2-0.33, 0.66, 0.66, RGBColor.from_string("1E2530"),
         line=CORAL, lw=1.4, shape=MSO_SHAPE.OVAL)                 # icon disc
    text(x+0.30, y+h/2-0.33, 0.66, 0.66, [[(icon, 20, CORAL, True)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(x+1.10, y+0.12, w-1.3, 0.35, [[(title, 13, MUTE, True)]])
    text(x+1.10, y+0.44, 2.4, 0.5, [[(big, 28, INK, True)]],
         anchor=MSO_ANCHOR.MIDDLE)
    text(x+w-2.3, y+0.50, 2.18, 0.4, [[(small, 12, MUTE, False)]],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

chip(0.95, "◎", "DICE SCORE  ·  MAE @10%", "0.86", "vs 0.93 @100%")
chip(2.18, "♥", "EF MAE  ·  MAE @10%",     "12.6%", "vs 6.9% @100%")
chip(3.41, "★", "LABEL EFFICIENCY",             "10×", "fewer labels")

# highlighted statement
rect(0.6, 4.95, 12.13, 0.90, HILITE, line=CORAL, lw=2.0,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(0.7, 4.95, 11.93, 0.90,
     [[("MAE-pretrained at 10% labels reaches 0.86 Dice — within 0.07 of the", 15, INK, True)],
      [("fully-supervised baseline trained on 100% of the labels.", 15, INK, True)]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# comparison table
rows, cols = 3, 4
tbl_shape = slide.shapes.add_table(rows, cols, Inches(0.6), Inches(6.05),
                                   Inches(12.13), Inches(1.05))
tbl = tbl_shape.table
# kill default banding style
tbl.first_row = False; tbl.horz_banding = False
for c, wdt in zip(range(cols), [4.6, 2.2, 3.2, 2.13]):
    tbl.columns[c].width = Inches(wdt)
data = [
    ["Method", "Labels", "Dice (2CH / 4CH)", "EF MAE"],
    ["MAE Pretrained", "10%", "0.862 / 0.864", "12.65%"],
    ["Supervised Baseline", "100%", "0.927 / 0.925", "6.86%"],
]
for r in range(rows):
    for c in range(cols):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = BG
        elif r == 1:
            cell.fill.fore_color.rgb = HILITE     # highlighted MAE row
        else:
            cell.fill.fore_color.rgb = BG
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.08); cell.margin_top = Inches(0.02)
        cell.margin_bottom = Inches(0.02)
        p = cell.text_frame.paragraphs[0]
        run = p.add_run(); run.text = data[r][c]
        run.font.name = FONT; run.font.size = Pt(12.5 if r == 0 else 13)
        run.font.bold = (r == 0) or (c == 0) or (r == 1)
        if r == 0:
            run.font.color.rgb = MUTE
        elif r == 1:
            run.font.color.rgb = INK
        else:
            run.font.color.rgb = MUTE

text(0.6, 7.15, 12.13, 0.3,
     [[("Source: CAMUS test set · 3 seeds · results/master_ablation_table.csv",
        9.5, RGBColor.from_string("5C6675"), False)]],
     align=PP_ALIGN.RIGHT)

prs.save("figures/slide13_headline_result.pptx")
print("Saved figures/slide13_headline_result.pptx")
